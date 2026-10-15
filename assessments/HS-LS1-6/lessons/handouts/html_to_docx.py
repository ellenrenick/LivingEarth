"""Convert a simple handout HTML file (h1, h2, p, b/strong, ol/ul > li, table) to .docx.

  python3 html_to_docx.py in.html out.docx [body_pt]
"""

import sys
from html.parser import HTMLParser

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt


class Builder(HTMLParser):
    def __init__(self, doc):
        super().__init__()
        self.doc, self.para, self.bold, self.lists = doc, None, 0, []
        self.table, self.row, self.cell_text, self.in_style = None, None, None, False

    def start_para(self, style=None):
        self.para = self.doc.add_paragraph(style=style)
        self.para.paragraph_format.space_after = Pt(3)

    def handle_starttag(self, tag, attrs):
        if tag == "style":
            self.in_style = True
        elif tag in ("h1", "h2"):
            self.start_para("Heading 1" if tag == "h1" else "Heading 2")
        elif tag == "p" and self.cell_text is None:
            self.start_para()
        elif tag in ("ol", "ul"):
            self.lists.append([tag, 0])
        elif tag == "li":
            kind = self.lists[-1]
            if kind[0] == "ol":
                # Number by hand so each list restarts at 1 (Word's List Number style keeps counting).
                kind[1] += 1
                self.start_para()
                self.para.paragraph_format.left_indent = Inches(0.3)
                self.para.paragraph_format.first_line_indent = Inches(-0.2)
                self.para.paragraph_format.space_after = Pt(1)
                self.para.add_run(f"{kind[1]}. ")
            else:
                self.start_para("List Bullet")
        elif tag in ("b", "strong"):
            self.bold += 1
        elif tag == "br" and self.para is not None:
            self.para.add_run().add_break()
        elif tag == "table":
            self.table = []
        elif tag == "tr":
            self.row = []
        elif tag in ("td", "th"):
            self.cell_text = [tag == "th", ""]

    def handle_endtag(self, tag):
        if tag == "style":
            self.in_style = False
        elif tag in ("h1", "h2", "p", "li") and self.cell_text is None:
            self.para = None
        elif tag in ("ol", "ul"):
            self.lists.pop()
        elif tag in ("b", "strong"):
            self.bold -= 1
        elif tag in ("td", "th"):
            self.row.append(self.cell_text)
            self.cell_text = None
        elif tag == "tr":
            self.table.append(self.row)
        elif tag == "table":
            rows = self.table
            t = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.LEFT
            for r, row in enumerate(rows):
                for c, (head, text) in enumerate(row):
                    run = t.cell(r, c).paragraphs[0].add_run(text.strip())
                    run.bold = head
            self.table = None

    def handle_data(self, data):
        if self.in_style:
            return
        if self.cell_text is not None:
            self.cell_text[1] += data
            return
        if self.para is None:
            if not data.strip():
                return
            self.start_para()
        text = " ".join(data.split())
        if data[:1].isspace() and self.para.runs:
            text = " " + text
        if data[-1:].isspace() and text:
            text += " "
        if text:
            self.para.add_run(text).bold = self.bold > 0


def main(src, out, body_pt="11"):
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.7)
        s.left_margin = s.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(float(body_pt))
    # Compact headings so a two-page plan stays on two pages.
    for name, size in (("Heading 1", 15), ("Heading 2", 12.5)):
        st = doc.styles[name]
        st.font.size = Pt(size)
        st.paragraph_format.space_before = Pt(8 if name == "Heading 2" else 0)
        st.paragraph_format.space_after = Pt(2)
    doc.styles["List Bullet"].paragraph_format.space_after = Pt(1)
    Builder(doc).feed(open(src, encoding="utf-8").read())
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main(*sys.argv[1:4])
