"""Render student-facing slide decks (.pptx) for the 8-day unit. Usage: python3 make_slides.py  (writes ./slides/Day*.pptx)"""
import re
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
import slides_8day

NAVY, TEAL, LIGHT, DARK = RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0x2A, 0x9D, 0x8F), RGBColor(0xEA, 0xF4, 0xF3), RGBColor(0x22, 0x22, 0x22)


def filename(title):
    return re.sub(r"[^A-Za-z0-9]+", "_", title.split(":")[0] + "_" + title.split(":", 1)[1]).strip("_")[:60] + ".pptx"


def add_slide(prs, s, n, total):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bar = sl.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.2))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    tb = sl.shapes.add_textbox(Inches(0.5), Inches(0.12), Inches(12.3), Inches(1.0))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = s["title"]; p.font.size = Pt(30 if len(s["title"]) < 55 else 26); p.font.bold = True; p.font.color.rgb = RGBColor(255, 255, 255)
    y = 1.5
    bullets, table = s.get("bullets", []), s.get("table")
    if bullets:
        h = 5.4 if not table else min(0.6 * len(bullets) + 0.3, 2.6)
        box = sl.shapes.add_textbox(Inches(0.7), Inches(y), Inches(11.9), Inches(h))
        tf = box.text_frame; tf.word_wrap = True
        for i, b in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = ("" if len(bullets) == 1 and not table and n == 1 and False else "• ") + b
            p.font.size = Pt(26 if not table else 20); p.font.color.rgb = DARK; p.space_after = Pt(10)
        y += h + 0.1
    if table:
        rows, cols = len(table["rows"]) + 1, len(table["header"])
        rh = 0.42 if rows > 6 else 0.6
        gt = sl.shapes.add_table(rows, cols, Inches(0.7), Inches(y), Inches(11.9), Inches(rh * rows)).table
        for c, hd in enumerate(table["header"]):
            cell = gt.cell(0, c); cell.text = hd; cell.fill.solid(); cell.fill.fore_color.rgb = TEAL
            for p in cell.text_frame.paragraphs: p.font.size = Pt(18 if rows <= 6 else 16); p.font.bold = True; p.font.color.rgb = RGBColor(255, 255, 255)
        for r, row in enumerate(table["rows"], 1):
            for c, v in enumerate(row):
                cell = gt.cell(r, c); cell.text = v; cell.fill.solid(); cell.fill.fore_color.rgb = LIGHT if r % 2 else RGBColor(255, 255, 255)
                for p in cell.text_frame.paragraphs: p.font.size = Pt(18 if rows <= 6 else 16); p.font.color.rgb = DARK
        y += rh * rows + 0.15
    if s.get("note"):
        nb = sl.shapes.add_textbox(Inches(0.7), Inches(min(y, 6.5)), Inches(11.9), Inches(0.5))
        nb.text_frame.word_wrap = True
        p = nb.text_frame.paragraphs[0]; p.text = s["note"]; p.font.size = Pt(16); p.font.italic = True; p.font.color.rgb = NAVY
    ft = sl.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(12.3), Inches(0.35))
    p = ft.text_frame.paragraphs[0]; p.text = f"Unit 2.1: DNA to Proteins  |  {n} of {total}"; p.font.size = Pt(11); p.font.color.rgb = RGBColor(0x66, 0x66, 0x66); p.alignment = PP_ALIGN.RIGHT


def main():
    out = Path(__file__).parent / "slides"; out.mkdir(exist_ok=True)
    for title, slides in slides_8day.DECKS.items():
        prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        for i, s in enumerate(slides, 1):
            add_slide(prs, s, i, len(slides))
        prs.save(out / filename(title)); print(filename(title), len(slides))


if __name__ == "__main__":
    main()
